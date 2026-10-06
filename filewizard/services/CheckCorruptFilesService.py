import codecs
import logging
import os
import shutil
import traceback
import zipfile
import zlib
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import olefile
from pyxlsb import open_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError


logger = logging.getLogger(__name__)


class CheckCorruptFilesService:

    FILE_TYPE_GROUPS = {
        'PDF': ('.pdf',),
        'Word documents': (
            '.doc',
            '.docx',
            '.docm',
            '.dot',
            '.dotx',
            '.rtf',
            '.odt',
        ),
        'PowerPoint': (
            '.ppt',
            '.pptx',
            '.pps',
            '.ppsx',
            '.pot',
            '.potx',
            '.odp',
        ),
        'Spreadsheets': (
            '.xls',
            '.xlsx',
            '.xlsm',
            '.xlsb',
            '.xlt',
            '.xltx',
            '.xltm',
            '.ods',
            '.csv',
            '.tsv',
        ),
        'Text': ('.txt',),
    }
    SUPPORTED_EXTENSIONS = tuple(
        extension
        for extensions in FILE_TYPE_GROUPS.values()
        for extension in extensions
    )
    EXTENSION_VALIDATORS = {
        '.pdf': 'pdf',
        '.doc': 'ole',
        '.docx': 'ooxml',
        '.docm': 'ooxml',
        '.dot': 'ole',
        '.dotx': 'ooxml',
        '.rtf': 'rtf',
        '.odt': 'opendocument',
        '.ppt': 'ole',
        '.pptx': 'ooxml',
        '.pps': 'ole',
        '.ppsx': 'ooxml',
        '.pot': 'ole',
        '.potx': 'ooxml',
        '.odp': 'opendocument',
        '.xls': 'ole',
        '.xlsx': 'ooxml',
        '.xlsm': 'ooxml',
        '.xlsb': 'xlsb',
        '.xlt': 'ole',
        '.xltx': 'ooxml',
        '.xltm': 'ooxml',
        '.ods': 'opendocument',
        '.csv': 'text',
        '.tsv': 'text',
        '.txt': 'text',
    }
    OLE_SIGNATURE = bytes.fromhex('D0CF11E0A1B11AE1')
    OOXML_PARTS = {
        'word': {
            '[Content_Types].xml': (
                'http://schemas.openxmlformats.org/package/2006/content-types',
                'Types',
            ),
            'word/document.xml': (
                'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                'document',
            ),
        },
        'powerpoint': {
            '[Content_Types].xml': (
                'http://schemas.openxmlformats.org/package/2006/content-types',
                'Types',
            ),
            'ppt/presentation.xml': (
                'http://schemas.openxmlformats.org/presentationml/2006/main',
                'presentation',
            ),
        },
        'excel': {
            '[Content_Types].xml': (
                'http://schemas.openxmlformats.org/package/2006/content-types',
                'Types',
            ),
            'xl/workbook.xml': (
                'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
                'workbook',
            ),
        },
    }
    OOXML_TYPES = {
        '.docx': 'word',
        '.docm': 'word',
        '.dotx': 'word',
        '.pptx': 'powerpoint',
        '.ppsx': 'powerpoint',
        '.potx': 'powerpoint',
        '.xlsx': 'excel',
        '.xlsm': 'excel',
        '.xltx': 'excel',
        '.xltm': 'excel',
    }
    OPEN_DOCUMENT_TYPES = {
        '.odt': 'application/vnd.oasis.opendocument.text',
        '.odp': 'application/vnd.oasis.opendocument.presentation',
        '.ods': 'application/vnd.oasis.opendocument.spreadsheet',
    }

    def __init__(
        self,
        folderPath,
        includeSubfolders=False,
        dryRun=True,
        progressCallback=None,
        fileExtensions=None,
    ):
        self.folderPath = os.path.abspath(folderPath)
        self.includeSubfolders = includeSubfolders
        self.dryRun = dryRun
        self.progressCallback = progressCallback
        self.fileExtensions = self._normalizeExtensions(fileExtensions)
        self.corruptFolder = os.path.join(self.folderPath, 'corrupt')

    def execute(self):
        self._validateFolder()
        files = self._getFiles()

        results = {
            'folderPath': self.folderPath,
            'total': len(files),
            'checked': [],
            'corrupt': [],
            'skipped': [],
            'failed': [],
            'dryRun': self.dryRun,
            'includeSubfolders': self.includeSubfolders,
            'fileExtensions': self.fileExtensions,
        }

        reservedDestinations = set()

        for current, filePath in enumerate(files, start=1):
            self._reportProgress(current, len(files), filePath)

            try:
                status, reason = self._validateFile(filePath)
            except Exception as exception:
                self._recordFailure(
                    results,
                    filePath,
                    'validation',
                    exception,
                )
                continue

            if status == 'skipped':
                results['skipped'].append({
                    'filePath': filePath,
                    'reason': reason,
                })
                continue

            if status == 'valid':
                results['checked'].append(filePath)
                continue

            destination = self._getAvailablePath(
                filePath,
                reservedDestinations,
            )
            item = {
                'filePath': filePath,
                'destination': destination,
                'reason': reason,
                'extension': os.path.splitext(filePath)[1].lower(),
            }

            if self.dryRun:
                results['corrupt'].append(item)
                continue

            try:
                os.makedirs(os.path.dirname(destination), exist_ok=True)
                shutil.move(filePath, destination)
                results['corrupt'].append(item)
            except OSError as exception:
                reservedDestinations.discard(
                    self._normalizePath(destination)
                )
                self._recordFailure(
                    results,
                    filePath,
                    f'move to {destination}',
                    exception,
                )

        return results

    def _validateFolder(self):
        if not os.path.isdir(self.folderPath):
            raise ValueError(
                f'Folder does not exist: {self.folderPath}'
            )

    def _normalizeExtensions(self, extensions):
        if extensions is None:
            return self.SUPPORTED_EXTENSIONS

        normalized = tuple(sorted({
            extension.lower()
            if extension.startswith('.')
            else f'.{extension.lower()}'
            for extension in extensions
        }))
        unsupported = set(normalized) - set(self.SUPPORTED_EXTENSIONS)

        if unsupported:
            raise ValueError(
                'Unsupported file extension(s): '
                + ', '.join(sorted(unsupported))
            )

        if not normalized:
            raise ValueError('Select at least one file type to check.')

        return normalized

    def _getFiles(self):
        files = []
        corruptFolder = self._normalizePath(self.corruptFolder)

        if self.includeSubfolders:
            for root, dirs, fileNames in os.walk(
                self.folderPath,
                onerror=self._raiseWalkError,
            ):
                dirs[:] = [
                    directory
                    for directory in dirs
                    if self._normalizePath(
                        os.path.join(root, directory)
                    ) != corruptFolder
                ]
                files.extend(
                    os.path.join(root, fileName)
                    for fileName in fileNames
                    if os.path.splitext(fileName)[1].lower()
                    in self.fileExtensions
                )
            return files

        with os.scandir(self.folderPath) as entries:
            for entry in entries:
                if (
                    entry.is_file()
                    and os.path.splitext(entry.name)[1].lower()
                    in self.fileExtensions
                ):
                    files.append(entry.path)

        return files

    def _raiseWalkError(self, exception):
        raise exception

    def _validateFile(self, filePath):
        extension = os.path.splitext(filePath)[1].lower()
        validator = self.EXTENSION_VALIDATORS.get(extension)
        if validator == 'pdf':
            return self._validatePdf(filePath)
        if validator == 'ole':
            return self._validateOleFile(filePath)
        if validator == 'ooxml':
            return self._validateOoxml(filePath, extension)
        if validator == 'opendocument':
            return self._validateOpenDocument(filePath, extension)
        if validator == 'xlsb':
            return self._validateXlsb(filePath)
        if validator == 'rtf':
            return self._validateRtf(filePath)
        if validator == 'text':
            return self._validateText(filePath)

        raise ValueError(f'No validator is available for {extension}.')

    def _validatePdf(self, filePath):
        try:
            reader = PdfReader(filePath, strict=False)
            if reader.is_encrypted:
                return (
                    'skipped',
                    'Encrypted PDFs cannot be inspected without a password.',
                )

            for page in reader.pages:
                page.mediabox
                contents = page.get_contents()
                if contents is not None:
                    contents.get_data()
        except NotImplementedError:
            return 'skipped', 'The PDF uses an encoding that cannot be inspected.'
        except (
            PdfReadError,
            EOFError,
            ValueError,
            KeyError,
            IndexError,
            zlib.error,
        ) as exception:
            return self._corruptReason(
                exception,
                'The PDF structure could not be read.',
            )

        return 'valid', ''

    def _validateOoxml(self, filePath, extension):
        documentType = self.OOXML_TYPES[extension]

        try:
            if self._hasOleSignature(filePath):
                return (
                    'skipped',
                    'Encrypted Office files cannot be inspected without a password.',
                )

            with zipfile.ZipFile(filePath) as archive:
                corruptEntry = archive.testzip()
                if corruptEntry is not None:
                    return 'corrupt', f'Invalid ZIP data in {corruptEntry}.'

                requiredParts = self.OOXML_PARTS[documentType]
                names = set(archive.namelist())
                missingFiles = [
                    name for name in requiredParts if name not in names
                ]
                if missingFiles:
                    return (
                        'corrupt',
                        'Missing required Office part(s): '
                        + ', '.join(missingFiles),
                    )

                for name, expectedRoot in requiredParts.items():
                    root = ElementTree.fromstring(archive.read(name))
                    expectedTag = (
                        f'{{{expectedRoot[0]}}}{expectedRoot[1]}'
                    )
                    if root.tag != expectedTag:
                        return (
                            'corrupt',
                            f'Invalid XML root in {name}.',
                        )

            if documentType == 'excel':
                workbook = load_workbook(
                    filePath,
                    read_only=True,
                    data_only=False,
                )
                try:
                    if not workbook.sheetnames:
                        return 'corrupt', 'The workbook contains no worksheets.'
                    for worksheet in workbook.worksheets:
                        for _ in worksheet.iter_rows():
                            pass
                finally:
                    workbook.close()
        except NotImplementedError:
            return 'skipped', 'The Office file uses an unsupported feature.'
        except RuntimeError as exception:
            return 'skipped', str(exception).strip() or 'The file is encrypted.'
        except (
            zipfile.BadZipFile,
            zipfile.LargeZipFile,
            ElementTree.ParseError,
            InvalidFileException,
            EOFError,
            ValueError,
            KeyError,
            IndexError,
            zlib.error,
        ) as exception:
            return self._corruptReason(
                exception,
                f'The {extension} package could not be read.',
            )

        return 'valid', ''

    def _validateXlsb(self, filePath):
        try:
            if self._hasOleSignature(filePath):
                return (
                    'skipped',
                    'Encrypted Office files cannot be inspected without a password.',
                )

            with zipfile.ZipFile(filePath) as archive:
                corruptEntry = archive.testzip()
                if corruptEntry is not None:
                    return 'corrupt', f'Invalid ZIP data in {corruptEntry}.'
                if 'xl/workbook.bin' not in archive.namelist():
                    return 'corrupt', 'Missing required part xl/workbook.bin.'

            with open_workbook(filePath) as workbook:
                sheets = workbook.sheets
                if not sheets:
                    return 'corrupt', 'The workbook contains no worksheets.'

                for sheetName in sheets:
                    with workbook.get_sheet(sheetName) as sheet:
                        for _ in sheet.rows():
                            break
        except NotImplementedError:
            return 'skipped', 'The XLSB file uses an unsupported feature.'
        except RuntimeError as exception:
            return 'skipped', str(exception).strip() or 'The file is encrypted.'
        except (
            zipfile.BadZipFile,
            EOFError,
            ValueError,
            KeyError,
            IndexError,
            zlib.error,
        ) as exception:
            return self._corruptReason(
                exception,
                'The XLSB workbook could not be read.',
            )

        return 'valid', ''

    def _validateOleFile(self, filePath):
        extension = os.path.splitext(filePath)[1].lower()
        requiredStreams = {
            '.doc': (('worddocument',),),
            '.dot': (('worddocument',),),
            '.ppt': (('powerpoint document',),),
            '.pps': (('powerpoint document',),),
            '.pot': (('powerpoint document',),),
            '.xls': (('workbook',), ('book',)),
            '.xlt': (('workbook',), ('book',)),
        }[extension]

        if not self._hasOleSignature(filePath):
            return 'corrupt', 'The file does not have a valid OLE signature.'

        try:
            compoundFile = olefile.OleFileIO(filePath)
            try:
                streams = {
                    tuple(part.casefold() for part in stream)
                    for stream in compoundFile.listdir()
                }
            finally:
                compoundFile.close()
        except (
            olefile.olefile.OleFileError,
            OSError,
            ValueError,
        ) as exception:
            return self._corruptReason(
                exception,
                'The legacy Office container could not be read.',
            )

        if not any(stream in streams for stream in requiredStreams):
            return (
                'corrupt',
                'The legacy Office file is missing its required document stream.',
            )

        return 'valid', ''

    def _hasOleSignature(self, filePath):
        with open(filePath, 'rb') as document:
            return (
                document.read(len(self.OLE_SIGNATURE))
                == self.OLE_SIGNATURE
            )

    def _validateOpenDocument(self, filePath, extension):
        expectedMimeType = self.OPEN_DOCUMENT_TYPES[extension]

        try:
            with zipfile.ZipFile(filePath) as archive:
                corruptEntry = archive.testzip()
                if corruptEntry is not None:
                    return 'corrupt', f'Invalid ZIP data in {corruptEntry}.'

                names = set(archive.namelist())
                if 'mimetype' not in names or 'content.xml' not in names:
                    return (
                        'corrupt',
                        'Missing required OpenDocument package parts.',
                    )

                if 'META-INF/manifest.xml' in names:
                    manifest = ElementTree.fromstring(
                        archive.read('META-INF/manifest.xml')
                    )
                    if any(
                        self._localName(element.tag) == 'encryption-data'
                        for element in manifest.iter()
                    ):
                        return (
                            'skipped',
                            'Encrypted OpenDocument files cannot be inspected.',
                        )

                mimeType = archive.read('mimetype').decode('ascii').strip()
                if mimeType != expectedMimeType:
                    return 'corrupt', 'The document has an invalid mimetype.'

                ElementTree.fromstring(archive.read('content.xml'))

                meta = 'meta.xml'
                if meta in names:
                    ElementTree.fromstring(archive.read(meta))
        except NotImplementedError:
            return 'skipped', 'The OpenDocument file uses an unsupported feature.'
        except (
            zipfile.BadZipFile,
            zipfile.LargeZipFile,
            ElementTree.ParseError,
            UnicodeDecodeError,
            EOFError,
            ValueError,
            KeyError,
            zlib.error,
        ) as exception:
            return self._corruptReason(
                exception,
                f'The {extension} package could not be read.',
            )

        return 'valid', ''

    def _validateRtf(self, filePath):
        with open(filePath, 'rb') as document:
            header = document.read(5)

        if not header.startswith(b'{\\rtf'):
            return 'corrupt', 'The file does not have a valid RTF header.'

        depth = 0
        escaped = False
        with open(filePath, 'rb') as document:
            while True:
                chunk = document.read(65536)
                if not chunk:
                    break
                for byte in chunk:
                    if escaped:
                        escaped = False
                    elif byte == ord('\\'):
                        escaped = True
                    elif byte == ord('{'):
                        depth += 1
                    elif byte == ord('}'):
                        depth -= 1
                        if depth < 0:
                            return 'corrupt', 'The RTF groups are unbalanced.'

        if depth != 0:
            return 'corrupt', 'The RTF groups are unbalanced.'

        return 'valid', ''

    def _localName(self, tag):
        return tag.rsplit('}', 1)[-1]

    def _validateText(self, filePath):
        encodings = (
            'utf-8-sig',
            'utf-32-le',
            'utf-32-be',
            'utf-16-le',
            'utf-16-be',
            'cp1252',
        )

        for encoding in encodings:
            decoder = codecs.getincrementaldecoder(encoding)(errors='strict')
            try:
                with open(filePath, 'rb') as textFile:
                    while True:
                        chunk = textFile.read(65536)
                        if not chunk:
                            break
                        decoder.decode(chunk)
                    decoder.decode(b'', final=True)
                return 'valid', ''
            except UnicodeError:
                continue

        return (
            'corrupt',
            'The file cannot be decoded as UTF-8, UTF-16, UTF-32, or Windows-1252.',
        )

    def _corruptReason(self, exception, default):
        reason = str(exception).strip()
        return 'corrupt', reason or default

    def _recordFailure(self, results, filePath, operation, exception):
        exceptionTraceback = traceback.format_exc()
        logger.error(
            'Failed during %s for file %s: %s: %s',
            operation,
            filePath,
            type(exception).__name__,
            exception,
            exc_info=True,
        )
        results['failed'].append({
            'filePath': filePath,
            'operation': operation,
            'errorType': type(exception).__name__,
            'reason': str(exception),
            'traceback': exceptionTraceback,
        })

    def _getAvailablePath(self, filePath, reservedDestinations):
        relativePath = os.path.relpath(filePath, self.folderPath)
        destination = os.path.join(self.corruptFolder, relativePath)

        if self._isAvailablePath(destination, reservedDestinations):
            reservedDestinations.add(self._normalizePath(destination))
            return destination

        baseName, extension = os.path.splitext(destination)
        counter = 1
        while True:
            candidate = f'{baseName}_{counter}{extension}'
            if self._isAvailablePath(candidate, reservedDestinations):
                reservedDestinations.add(self._normalizePath(candidate))
                return candidate
            counter += 1

    def _isAvailablePath(self, path, reservedDestinations):
        return (
            not os.path.lexists(path)
            and self._normalizePath(path) not in reservedDestinations
        )

    def _normalizePath(self, path):
        return os.path.normcase(os.path.abspath(path))

    def _reportProgress(self, current, total, filePath):
        if self.progressCallback:
            self.progressCallback(current, total, filePath)
