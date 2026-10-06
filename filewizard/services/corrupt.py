import os
import shutil
import zipfile
import logging

from openpyxl import load_workbook


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

logger = logging.getLogger(__name__)


# ============================================================
# FILE CORRUPTION CHECK FUNCTIONS
# ============================================================

def is_xlsx_corrupt(file_path):
    """
    Returns:
        True  = corrupt
        False = OK
    """
    try:
        if not zipfile.is_zipfile(file_path):
            return True

        wb = load_workbook(
            file_path,
            read_only=True,
            data_only=False
        )

        _ = wb.sheetnames

        wb.close()

        return False

    except Exception:
        return True


def is_pdf_corrupt(file_path):
    """
    Check whether a PDF file is corrupt.
    """
    try:
        with open(file_path, "rb") as f:
            header = f.read(5)

        if header != b"%PDF-":
            return True

        from pypdf import PdfReader

        reader = PdfReader(file_path)

        _ = len(reader.pages)

        return False

    except Exception:
        return True


def is_txt_corrupt(file_path):
    """
    Check whether a TXT file is readable.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            f.read()

        return False

    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="cp1252") as f:
                f.read()

            return False

        except Exception:
            return True

    except Exception:
        return True


def is_docx_corrupt(file_path):
    """
    Check whether a DOCX file is corrupt.
    """
    try:
        if not zipfile.is_zipfile(file_path):
            return True

        with zipfile.ZipFile(file_path, "r") as z:

            bad_file = z.testzip()

            if bad_file is not None:
                return True

            required_files = [
                "[Content_Types].xml",
                "word/document.xml",
            ]

            for required_file in required_files:
                if required_file not in z.namelist():
                    return True

        return False

    except Exception:
        return True


def is_doc_corrupt(file_path):
    """
    Basic validation of a legacy .doc file.
    """
    try:
        with open(file_path, "rb") as f:
            header = f.read(8)

        expected_header = bytes.fromhex(
            "D0CF11E0A1B11AE1"
        )

        if header != expected_header:
            return True

        return False

    except Exception:
        return True


def is_xlsb_corrupt(file_path):
    """
    Check whether an XLSB file is corrupt.

    Returns:
        True  = corrupt
        False = OK
    """
    try:
        from pyxlsb import open_workbook

        with open_workbook(file_path) as wb:
            # Force pyxlsb to read the workbook structure
            sheets = wb.sheets

            # Make sure the workbook contains at least one sheet
            if not sheets:
                return True

            # Try opening each worksheet.
            # This forces additional parsing of the XLSB structure.
            for sheet_name in sheets:
                with wb.get_sheet(sheet_name) as sheet:
                    # Read the first row if one exists.
                    # This causes pyxlsb to parse worksheet data.
                    for _ in sheet.rows():
                        break

        return False

    except Exception:
        return True


# ============================================================
# FILE EXTENSION -> CORRUPTION CHECK FUNCTION
# ============================================================

CORRUPT_CHECKS = {
    ".xlsx": is_xlsx_corrupt,
    ".pdf": is_pdf_corrupt,
    ".txt": is_txt_corrupt,
    ".doc": is_doc_corrupt,
    ".docx": is_docx_corrupt,
    ".xlsb": is_xlsb_corrupt,
}


# ============================================================
# CHECK ALL FILES
# ============================================================

def check_files(folder_path):

    if not os.path.isdir(folder_path):
        logger.error(
            "Folder does not exist: %s",
            folder_path
        )
        return False

    found_corrupt = False

    # Keep track of extensions we've already logged.
    # This prevents the same unsupported extension from
    # being logged thousands of times.
    unsupported_extensions = set()

    for root, dirs, files in os.walk(folder_path):

        # Don't enter corrupted_file folders
        dirs[:] = [
            d for d in dirs
            if d.lower() != "corrupted_file"
        ]

        for filename in files:

            file_path = os.path.join(
                root,
                filename
            )

            # Get the file extension
            extension = os.path.splitext(
                filename
            )[1].lower()

            # ==================================================
            # FILE TYPE NOT IMPLEMENTED
            # ==================================================

            if extension not in CORRUPT_CHECKS:

                # Only log each unsupported extension once
                if extension not in unsupported_extensions:

                    unsupported_extensions.add(extension)

                    logger.warning(
                        "No corruption check implemented for "
                        "file type '%s'. Example file: %s",
                        extension if extension else "[NO EXTENSION]",
                        file_path
                    )

                continue

            # ==================================================
            # FILE TYPE IS IMPLEMENTED
            # ==================================================

            check_function = CORRUPT_CHECKS[extension]

            logger.info(
                "Checking: %s",
                file_path
            )

            try:
                is_corrupt = check_function(file_path)

            except Exception as e:

                # If the corruption-check function itself
                # unexpectedly fails, treat the file as corrupt
                # and log the error.
                logger.exception(
                    "Error checking file: %s",
                    file_path
                )

                is_corrupt = True

            # ==================================================
            # CORRUPT FILE
            # ==================================================

            if is_corrupt:

                found_corrupt = True

                logger.error(
                    "CORRUPT FILE FOUND: %s",
                    file_path
                )

                # Create corrupted_file folder
                # in the SAME directory as the original file.
                corrupted_folder = os.path.join(
                    root,
                    "corrupted_file"
                )

                os.makedirs(
                    corrupted_folder,
                    exist_ok=True
                )

                # Destination path
                destination = os.path.join(
                    corrupted_folder,
                    filename
                )

                # ==================================================
                # PREVENT OVERWRITING EXISTING FILES
                # ==================================================

                if os.path.exists(destination):

                    base, extension_with_dot = os.path.splitext(
                        filename
                    )

                    counter = 1

                    while os.path.exists(destination):

                        new_filename = (
                            f"{base}_{counter}"
                            f"{extension_with_dot}"
                        )

                        destination = os.path.join(
                            corrupted_folder,
                            new_filename
                        )

                        counter += 1

                # ==================================================
                # MOVE CORRUPT FILE
                # ==================================================

                try:

                    shutil.move(
                        file_path,
                        destination
                    )

                    logger.error(
                        "Moved corrupt file to: %s",
                        destination
                    )

                except Exception:

                    logger.exception(
                        "Failed to move corrupt file: %s",
                        file_path
                    )

            else:

                logger.info(
                    "OK: %s",
                    file_path
                )

    # ============================================================
    # LOG SUMMARY OF UNSUPPORTED FILE TYPES
    # ============================================================

    if unsupported_extensions:

        logger.warning(
            "The following file types do not have "
            "corruption checks implemented:"
        )

        for extension in sorted(unsupported_extensions):

            logger.warning(
                "  %s",
                extension if extension else "[NO EXTENSION]"
            )

    return found_corrupt


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    # Change this to the folder you want to scan
    folder = r"D:\\Photos & Videos\\Sorting\\Mufina backup\\files\\mov"

    result = check_files(folder)

    print()
    print("=" * 70)

    if result:
        print("RESULT: Corrupt files were found.")
        print("RESULT VALUE:", True)
    else:
        print("RESULT: No corrupt files were found.")
        print("RESULT VALUE:", False)

    print("=" * 70)
